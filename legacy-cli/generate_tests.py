#!/usr/bin/env python3
"""Generate test cases from automation/QA scripts using an LLM (OpenRouter or Ollama)."""
import argparse
import json
import os
import re
import sys
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

EXTENSIONS = {".py", ".js", ".ts", ".java", ".robot", ".cs", ".rb", ".ps1", ".sh"}
PROVIDERS = {
    "openrouter": {
        "url": "https://openrouter.ai/api/v1/chat/completions",
        "default_model": "openai/gpt-4o-mini",
    },
    "ollama": {
        "url": "http://localhost:11434/v1/chat/completions",
        "default_model": "llama3.1",
    },
}
FIELDS = ["id", "title", "type", "priority", "preconditions", "steps", "expected_result"]
CHUNK_CHARS = 12000

PROMPT = """You are a senior QA engineer. Read the automation/QA script below and write test cases
that cover what it does: positive flows, negative cases and edge cases.

Return ONLY a JSON array (no prose, no markdown fences). Each item must have keys:
"id" (TC-001 style), "title", "type" (positive|negative|edge), "priority" (High|Medium|Low),
"preconditions" (string), "steps" (array of strings), "expected_result" (string).

Script file: {name}

{code}
"""


def call_llm(provider, model, prompt, timeout):
    cfg = PROVIDERS[provider]
    headers = {"Content-Type": "application/json"}
    if provider == "openrouter":
        key = os.environ.get("OPENROUTER_API_KEY")
        if not key:
            sys.exit("Error: set OPENROUTER_API_KEY environment variable.")
        headers["Authorization"] = f"Bearer {key}"
    body = {
        "model": model or cfg["default_model"],
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.2,
    }
    try:
        r = requests.post(cfg["url"], headers=headers, json=body, timeout=timeout)
        r.raise_for_status()
    except requests.ConnectionError:
        sys.exit(f"Error: cannot reach {cfg['url']}. Is {provider} running/reachable?")
    except requests.RequestException as e:
        sys.exit(f"Error calling {provider}: {e}")
    return r.json()["choices"][0]["message"]["content"]


def parse_json(text):
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\[.*\]", text, re.S)
        if not m:
            raise
        data = json.loads(m.group(0))
    if not isinstance(data, list):
        raise ValueError("expected a JSON array")
    return data


def chunks(code):
    if len(code) <= CHUNK_CHARS:
        return [code]
    parts, cur = [], ""
    for line in code.splitlines(keepends=True):
        if len(cur) + len(line) > CHUNK_CHARS and cur:
            parts.append(cur)
            cur = ""
        cur += line
    if cur:
        parts.append(cur)
    return parts


def generate_for_file(path, args):
    code = path.read_text(encoding="utf-8", errors="replace")
    cases = []
    for part in chunks(code):
        prompt = PROMPT.format(name=path.name, code=part)
        for attempt in range(1, args.retries + 1):
            try:
                cases += parse_json(call_llm(args.provider, args.model, prompt, args.timeout))
                break
            except (json.JSONDecodeError, ValueError, KeyError) as e:
                print(f"  bad LLM output (attempt {attempt}/{args.retries}): {e}")
        else:
            print(f"  skipped a chunk of {path.name}: no valid JSON")
    for i, c in enumerate(cases, 1):  # renumber so ids are unique across chunks
        c["id"] = f"TC-{i:03d}"
    return cases


def as_text(v):
    if isinstance(v, list):
        return "\n".join(f"{i}. {s}" for i, s in enumerate(v, 1))
    return str(v or "")


def write_markdown(path, cases, out_dir):
    lines = [f"# Test cases: {path.name}", "", f"Total: {len(cases)}", ""]
    lines += ["| ID | Title | Type | Priority |", "|---|---|---|---|"]
    for c in cases:
        lines.append(f"| {c.get('id','')} | {c.get('title','')} | {c.get('type','')} | {c.get('priority','')} |")
    for c in cases:
        lines += ["", f"## {c.get('id','')} - {c.get('title','')}",
                  f"**Type:** {c.get('type','')}  |  **Priority:** {c.get('priority','')}", "",
                  f"**Preconditions:** {c.get('preconditions','') or 'None'}", "", "**Steps:**", ""]
        steps = c.get("steps", [])
        lines += [f"{i}. {s}" for i, s in enumerate(steps if isinstance(steps, list) else [steps], 1)]
        lines += ["", f"**Expected result:** {c.get('expected_result','')}"]
    (out_dir / f"{path.stem}_testcases.md").write_text("\n".join(lines), encoding="utf-8")
    (out_dir / f"{path.stem}_testcases.json").write_text(json.dumps(cases, indent=2), encoding="utf-8")


def write_excel(all_cases, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(["Script"] + [f.replace("_", " ").title() for f in FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for script, c in all_cases:
        ws.append([script] + [as_text(c.get(f)) for f in FIELDS])
    for col, w in zip("ABCDEFGH", [22, 10, 40, 11, 10, 30, 50, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_dir / "test_cases.xlsx")


def ask(question, default=None):
    """Prompt the user; Enter accepts the default."""
    suffix = f" [{default}]" if default else ""
    answer = input(f"{question}{suffix}: ").strip().strip('"')
    return answer or default


def interactive(args):
    """Guided mode: asks simple questions when the tool is run with no options."""
    print("=== Test Case Generator ===")
    print("Answer a few questions (press Enter to accept the [default]).\n")
    args.input = ask("1. Path to your script or folder")
    choice = ask("2. Which LLM? (1 = Ollama local, 2 = OpenRouter cloud)", "1")
    args.provider = "openrouter" if choice == "2" else "ollama"
    args.model = ask("3. Model name", PROVIDERS[args.provider]["default_model"])
    if args.provider == "openrouter" and not os.environ.get("OPENROUTER_API_KEY"):
        os.environ["OPENROUTER_API_KEY"] = ask("4. Paste your OpenRouter API key")
    print()
    return args


def main():
    ap = argparse.ArgumentParser(
        description="Generate test cases from automation/QA scripts using an LLM. "
                    "Run with no options for a guided mode.")
    ap.add_argument("--input", help="script file or folder")
    ap.add_argument("--output", default="output", help="output folder (default: output)")
    ap.add_argument("--provider", choices=PROVIDERS, default="ollama")
    ap.add_argument("--model", help="model name (provider default if omitted)")
    ap.add_argument("--retries", type=int, default=3)
    ap.add_argument("--timeout", type=int, default=300)
    args = ap.parse_args()
    if not args.input:
        args = interactive(args)

    src = Path(args.input or "")
    if not args.input or not src.exists():
        sys.exit(f"Error: path '{args.input}' not found. Check the path and try again.")
    files = [src] if src.is_file() else sorted(
        p for p in src.rglob("*") if p.is_file() and p.suffix in EXTENSIONS)
    if not files:
        sys.exit("No script files found (supported: " + " ".join(sorted(EXTENSIONS)) + ").")
    out_dir = Path(args.output)
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Found {len(files)} script(s). Using {args.provider}. This may take a minute...\n")
    all_cases = []
    for n, f in enumerate(files, 1):
        print(f"[{n}/{len(files)}] {f.name} ...")
        cases = generate_for_file(f, args)
        if cases:
            write_markdown(f, cases, out_dir)
            all_cases += [(f.name, c) for c in cases]
        print(f"      -> {len(cases)} test cases")
    if all_cases:
        write_excel(all_cases, out_dir)
    print(f"\nDone! {len(all_cases)} test cases saved in: {out_dir.resolve()}")
    print("Open test_cases.xlsx (all cases) or the .md files (one per script).")
    if all_cases and sys.platform == "win32":
        os.startfile(out_dir.resolve())


if __name__ == "__main__":
    main()
